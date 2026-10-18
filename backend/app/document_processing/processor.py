"""Conservative, provenance-tracked extraction; OCR is optional when Tesseract exists."""
from __future__ import annotations
from pathlib import Path
import hashlib, io, re, zipfile
import fitz
from PIL import Image, ImageOps, UnidentifiedImageError
from docx import Document as DocxDocument
from app.core.config import MAX_PDF_PAGES, MAX_IMAGE_PIXELS

Image.MAX_IMAGE_PIXELS=MAX_IMAGE_PIXELS

MIMES = {'.pdf':'application/pdf', '.png':'image/png','.jpg':'image/jpeg','.jpeg':'image/jpeg', '.docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}
PATTERNS = {
 'policy_number':r'\bpolicy\s*(?:number|no\.?|#)\s*[:=#-]?\s*([A-Za-z0-9-]{4,35})',
 'invoice_number':r'\binvoice\s*(?:number|no\.?|#)\s*[:=#-]?\s*([A-Za-z0-9-]{3,35})',
 'claim_amount':r'\b(?:claim\s*amount|amount\s*claimed)\s*[:=]?\s*(?:INR|Rs\.?|₹|\$)?\s*([\d,]+(?:\.\d{1,2})?)',
 'invoice_total':r'\b(?:invoice\s*total|bill\s*total|total\s*invoice)\s*[:=]?\s*(?:INR|Rs\.?|₹|\$)?\s*([\d,]+(?:\.\d{1,2})?)',
 'incident_date':r'\b(?:incident|accident|loss)\s*date\s*[:=]?\s*(\d{4}-\d{2}-\d{2})',
 'policy_start':r'\b(?:policy\s*start|coverage\s*start)\s*[:=]?\s*(\d{4}-\d{2}-\d{2})',
 'policy_end':r'\b(?:policy\s*end|coverage\s*end)\s*[:=]?\s*(\d{4}-\d{2}-\d{2})',
 'vehicle_registration':r'\b(?:vehicle\s*(?:registration|reg)|registration\s*(?:number|no))\s*[:=#-]?\s*([A-Za-z0-9 -]{5,20})',
 'reported_date':r'\breported\s*date\s*[:=]?\s*(\d{4}-\d{2}-\d{2})',
 'coverage_amount':r'\bcoverage\s*amount\s*[:=]?\s*(?:INR|Rs\.?|₹|\$)?\s*([\d,]+(?:\.\d{1,2})?)',
 'prior_claims':r'\bprior\s*claims\s*[:=]?\s*(\d+)',
 'hospital_name':r'\bhospital\s*name\s*[:=]?\s*([^\r\n]{3,100})',
 'admission_date':r'\badmission\s*date\s*[:=]?\s*(\d{4}-\d{2}-\d{2})',
 'discharge_date':r'\bdischarge\s*date\s*[:=]?\s*(\d{4}-\d{2}-\d{2})',
 'trip_start':r'\btrip\s*start\s*[:=]?\s*(\d{4}-\d{2}-\d{2})',
 'trip_end':r'\btrip\s*end\s*[:=]?\s*(\d{4}-\d{2}-\d{2})',
 'booking_reference':r'\bbooking\s*reference\s*[:=#-]?\s*([A-Za-z0-9-]{4,50})',
 'repair_estimate':r'\brepair\s*estimate\s*[:=]?\s*(?:INR|Rs\.?|₹|\$)?\s*([\d,]+(?:\.\d{1,2})?)',
 'vehicle_type':r'\bvehicle\s*type\s*[:=]?\s*([^\r\n]{3,60})',
 'insurance_category':r'\binsurance\s*category\s*[:=]?\s*([A-Za-z _-]{3,40})',
}

def _validate_magic(data, suffix):
    if suffix=='.pdf' and not data.startswith(b'%PDF-'): raise ValueError('Invalid PDF file signature')
    if suffix=='.png' and not data.startswith(b'\x89PNG\r\n\x1a\n'): raise ValueError('Invalid PNG file signature')
    if suffix in ('.jpg','.jpeg') and not data.startswith(b'\xff\xd8\xff'): raise ValueError('Invalid JPEG file signature')
    if suffix=='.docx':
        if not zipfile.is_zipfile(io.BytesIO(data)): raise ValueError('Invalid DOCX file')
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if 'word/document.xml' not in z.namelist(): raise ValueError('Invalid DOCX structure')
            if sum(i.file_size for i in z.infolist()) > 30_000_000: raise ValueError('DOCX uncompressed size too large')
            if any(i.filename.startswith('../') for i in z.infolist()): raise ValueError('Invalid DOCX path')

def extract_fields(pages, doc_id):
    found={}
    for page_number,text in pages:
        for key,pattern in PATTERNS.items():
            m=re.search(pattern,text,re.IGNORECASE)
            if not m: continue
            value=m.group(1).strip()
            if key in ('claim_amount','invoice_total','coverage_amount','repair_estimate'):
                try: value=float(value.replace(',',''))
                except ValueError: continue
            if key=='prior_claims': value=int(value)
            if key not in found:
                found[key]={'value':value,'source_document':doc_id,'page':page_number,
                            'span':[m.start(1),m.end(1)],'method':'regex',
                            'confidence':None,'validated':False,'human_corrected':False}
    return found

def _ocr(image):
    try:
        import pytesseract
        return pytesseract.image_to_string(image)
    except (ImportError, OSError, RuntimeError) as exc:
        raise RuntimeError('OCR unavailable: install Tesseract binary and pytesseract') from exc
    except Exception as exc:
        if 'tesseract' in str(exc).lower(): raise RuntimeError('OCR unavailable: install Tesseract binary') from exc
        raise

def dhash(image):
    gray=ImageOps.grayscale(image).resize((9,8))
    pixels=list(gray.get_flattened_data())
    bits=0
    for y in range(8):
        for x in range(8): bits=(bits<<1)|int(pixels[y*9+x]>pixels[y*9+x+1])
    return f'{bits:016x}'

def hamming(hash_a,hash_b):
    return (int(hash_a,16)^int(hash_b,16)).bit_count()

def assess_image(data):
    """Technical observations only. Not a manipulation or fraud classifier."""
    import cv2, numpy as np
    with Image.open(io.BytesIO(data)) as img:
        img.verify()
    with Image.open(io.BytesIO(data)) as im:
        original=ImageOps.exif_transpose(im)
        if original.width*original.height>MAX_IMAGE_PIXELS: raise ValueError('Image exceeds pixel limit')
        rgb=original.convert('RGB')
        grayscale=np.asarray(rgb.convert('L'))
        variance=float(cv2.Laplacian(grayscale,cv2.CV_64F).var())
        return {'width':rgb.width,'height':rgb.height,'blur_indicator':round(variance,2),
                'is_low_resolution':min(rgb.size)<400, 'dhash':dhash(rgb),
                'note':'Quality measures cannot establish fraud or alteration.'}

def process_document(data:bytes, filename:str, doc_id:str):
    suffix=Path(filename).suffix.lower()
    if suffix not in MIMES: raise ValueError('Only PDF, PNG, JPG, JPEG and DOCX are supported as evidence')
    _validate_magic(data,suffix)
    pages=[]; warnings=[]; assessment={}; digest=hashlib.sha256(data).hexdigest()
    if suffix=='.pdf':
        try:
            with fitz.open(stream=data,filetype='pdf') as pdf:
                if pdf.needs_pass: raise ValueError('Encrypted PDFs are not supported')
                if len(pdf)>MAX_PDF_PAGES: raise ValueError(f'PDF exceeds {MAX_PDF_PAGES} pages')
                for i,page in enumerate(pdf):
                    text=page.get_text('text')
                    if len(text.strip())<20:
                        try:
                            pix=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False)
                            img=Image.open(io.BytesIO(pix.tobytes('png')))
                            text+='\n'+_ocr(img)
                        except RuntimeError as e: warnings.append(f'Page {i+1}: {e}')
                    pages.append((i+1,text[:150000]))
        except fitz.FileDataError as e: raise ValueError('Invalid or corrupted PDF') from e
    elif suffix=='.docx':
        try:
            doc=DocxDocument(io.BytesIO(data))
            text='\n'.join([p.text for p in doc.paragraphs]+[cell.text for table in doc.tables for row in table.rows for cell in row.cells])
            pages=[(1,text[:150000])]
        except (ValueError,KeyError,zipfile.BadZipFile) as e: raise ValueError('Invalid DOCX') from e
    else:
        try:
            assessment=assess_image(data)
            with Image.open(io.BytesIO(data)) as img:
                try: pages=[(1,_ocr(ImageOps.exif_transpose(img).convert('RGB'))) ]
                except RuntimeError as e: warnings.append(str(e)); pages=[(1,'')]
        except (UnidentifiedImageError,Image.DecompressionBombError) as e: raise ValueError('Invalid or oversized image') from e
    fields=extract_fields(pages,doc_id)
    return {'mime':MIMES[suffix],'sha256':digest,'dhash':assessment.get('dhash'),
            'extracted':fields,'warnings':warnings,'image_assessment':assessment,
            'pages_processed':len(pages)}
